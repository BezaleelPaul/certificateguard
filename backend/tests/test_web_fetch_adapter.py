"""Tests for the real-world HTTP verification adapter and adapter-chain gating.

These tests never touch the network: `_afetch` is stubbed and localhost URLs
plus `allow_localhost` metadata keep the SSRF validator deterministic offline.
"""

import base64
import io
import json

import httpx
import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.batch.analyzer import _resolve_issuer_record
from app.batch.template import (
    ALL_COLUMNS,
    INPUT_COLUMNS,
    SHEET_NAME,
    VERDICT_ANOMALY,
    VERDICT_FAKE,
    VERDICT_LEGIT,
    parse_workbook,
)
from app.main import app
from app.models import Issuer, IssuerVerificationStatus
from app.verification.adapter import AdapterVerificationResult
from app.verification.mock_adapter import mock_issuer_adapter
from app.verification.playwright_adapter import playwright_issuer_adapter
from app.verification.web_fetch_adapter import web_fetch_issuer_adapter

WORKBOOK_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

LOCAL_URL = "http://localhost:8001/certificates/MGL-001"
LOCAL_META = {"allow_localhost": True}

MGL_HTML = """<!doctype html><html><head>
<title>Programming Fundamentals course completion certificate for A Chirag Kevin Bernard</title>
<meta name="description" content="The certificate verifies that A Chirag Kevin Bernard has successfully completed the free course Programming Fundamentals.">
</head><body></body></html>"""


def _resp(url, status=200, text="", content_type="text/html; charset=utf-8"):
    return httpx.Response(
        status,
        headers={"content-type": content_type},
        text=text,
        request=httpx.Request("GET", url),
    )


def _install_fetch(monkeypatch, responder):
    async def fake_afetch(url: str) -> httpx.Response:
        return responder(url)

    monkeypatch.setattr(web_fetch_issuer_adapter, "_afetch", fake_afetch)


def _install_fetch_boom(monkeypatch):
    async def boom(url: str):
        raise AssertionError(f"fetch must not happen, got {url}")

    monkeypatch.setattr(web_fetch_issuer_adapter, "_afetch", boom)


# ---------------------------------------------------------------------------
# Adapter unit tests
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_extracts_recipient_from_title_and_meta(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, MGL_HTML))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned == "A Chirag Kevin Bernard"
    assert result.course_returned == "Programming Fundamentals"
    assert result.certificate_id_returned == "MGL-001"


@pytest.mark.asyncio
async def test_extracts_recipient_from_redirect_token_payload(monkeypatch):
    token = base64.urlsafe_b64encode(
        json.dumps(
            {
                "course_id": "1088",
                "certificate_url": "https://certificates.example.net/x.png",
                "username": "Abhishek K",
            }
        ).encode()
    ).decode()
    landing = f"http://localhost:8001/skillup-certificate-landing?token={token}"
    _install_fetch(
        monkeypatch, lambda url: _resp(url, 200, "<html><body></body></html>")
    )

    result = await web_fetch_issuer_adapter.verify(
        certificate_id="SK-77",
        verification_url=landing,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned == "Abhishek K"
    assert result.course_returned is None or isinstance(result.course_returned, str)


@pytest.mark.asyncio
async def test_http_404_is_invalid(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 404, "<html>nope</html>"))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-404",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.INVALID
    assert "HTTP 404" in (result.error_message or "")


@pytest.mark.asyncio
async def test_http_403_is_unavailable_not_invalid(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 403, "forbidden"))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-403",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "403" in (result.error_message or "")


@pytest.mark.asyncio
async def test_soft_404_page_is_invalid(monkeypatch):
    body = "<html><body><p>This certificate was not found in our records.</p></body></html>"
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, body))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-000",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.INVALID


@pytest.mark.asyncio
async def test_200_without_certificate_data_is_unavailable(monkeypatch):
    body = "<html><head><title>Acme Learning</title></head><body>Welcome</body></html>"
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, body))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE


@pytest.mark.asyncio
async def test_missing_url_is_unavailable_without_fetching(monkeypatch):
    _install_fetch_boom(monkeypatch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url=None,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "no public certificate URL" in (result.error_message or "")


@pytest.mark.asyncio
async def test_non_html_content_is_unavailable(monkeypatch):
    _install_fetch(
        monkeypatch,
        lambda url: _resp(url, 200, "binary", content_type="image/png"),
    )
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE


@pytest.mark.asyncio
async def test_ssrf_blocked_target_is_invalid_without_fetching(monkeypatch):
    _install_fetch_boom(monkeypatch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url="http://127.0.0.1/steal-metadata",
        metadata=None,
    )
    assert result.verification_status == IssuerVerificationStatus.INVALID
    assert "SSRF" in (result.error_message or "")


@pytest.mark.asyncio
async def test_single_word_title_never_becomes_a_recipient(monkeypatch):
    body = "<html><head><title>Certificate for free</title></head><body></body></html>"
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, body))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert result.recipient_returned is None


SKILLUP_OG_HTML = """<html><head>
<meta property="og:title" content="BHAGIRATHI  has successfully completed the online course Fundamentals of Software Development. Upgrade your skills with 100+ free courses">
<title>Skillup Certificate | Simplilearn</title>
</head><body></body></html>"""


@pytest.mark.asyncio
async def test_extracts_recipient_and_course_from_og_title(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, SKILLUP_OG_HTML))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="25BBTCS040",
        verification_url="http://localhost:8001/landing",
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned == "BHAGIRATHI"
    assert result.course_returned == "Fundamentals of Software Development"


@pytest.mark.asyncio
async def test_og_meta_content_before_property_order(monkeypatch):
    body = (
        '<html><head><meta content="Anish Jaiswal has successfully completed '
        'the online course Programming Fundamentals." property="og:title">'
        "</head><body></body></html>"
    )
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, body))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="25BBTCS055",
        verification_url="http://localhost:8001/landing",
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned == "Anish Jaiswal"


@pytest.mark.asyncio
async def test_verification_works_without_certificate_id(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 200, MGL_HTML))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="",
        verification_url=LOCAL_URL,
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned == "A Chirag Kevin Bernard"


@pytest.mark.asyncio
async def test_newline_inside_url_is_repaired_before_fetch(monkeypatch):
    seen = {}

    async def fake_afetch(url: str) -> httpx.Response:
        seen["url"] = url
        return _resp(url, 200, MGL_HTML)

    monkeypatch.setattr(web_fetch_issuer_adapter, "_afetch", fake_afetch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url="http://localhost:8001/certificates/\nmgl001",
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert seen["url"] == "http://localhost:8001/certificates/mgl001"


@pytest.mark.asyncio
async def test_two_links_in_one_cell_use_the_first(monkeypatch):
    seen = {}

    async def fake_afetch(url: str) -> httpx.Response:
        seen["url"] = url
        return _resp(url, 200, MGL_HTML)

    monkeypatch.setattr(web_fetch_issuer_adapter, "_afetch", fake_afetch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url="http://localhost:8001/one\nhttp://localhost:8001/two",
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert seen["url"] == "http://localhost:8001/one"


@pytest.mark.asyncio
async def test_non_http_url_text_is_unavailable_without_fetching(monkeypatch):
    _install_fetch_boom(monkeypatch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="MGL-001",
        verification_url="see the portal for details",
        metadata=LOCAL_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "not a fetchable http(s) link" in (result.error_message or "")


def test_parse_workbook_repairs_wrapped_url_cells():
    rows = [row for row in MGL_ROWS]
    rows[0] = list(rows[0])
    rows[0][4] = "https://www.mygreatlearning.com/certificates/\nmgl001"
    parsed = parse_workbook(_make_workbook(rows))
    assert "\n" not in parsed[0].certificate_url
    assert (
        parsed[0].certificate_url
        == "https://www.mygreatlearning.com/certificates/mgl001"
    )


# ---------------------------------------------------------------------------
# Analyzer adapter-chain gating
# ---------------------------------------------------------------------------


def _web_result(status, **kwargs):
    return AdapterVerificationResult(verification_status=status, **kwargs)


async def _mock_boom(*args, **kwargs):
    raise AssertionError("mock registry queried for a non-demo issuer")


async def _playwright_boom(*args, **kwargs):
    raise AssertionError("playwright queried for a non-demo issuer")


@pytest.mark.asyncio
async def test_real_issuer_uses_web_fetch_and_never_mock(monkeypatch):
    monkeypatch.setattr(mock_issuer_adapter, "verify", _mock_boom)
    monkeypatch.setattr(playwright_issuer_adapter, "verify", _playwright_boom)

    async def fake_web(**kwargs):
        return _web_result(
            IssuerVerificationStatus.UNAVAILABLE,
            error_message="Issuer site returned HTTP 403 (access restricted or outage)",
        )

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", fake_web)

    issuer = Issuer(
        name="MyGreatLearning",
        official_domain="mygreatlearning.com",
        verification_type="WEB",
        active=True,
    )
    result = await _resolve_issuer_record(
        issuer, "25BBTCS001", "https://www.mygreatlearning.com/certificate/x"
    )
    assert result["status"] == "UNAVAILABLE"
    assert result["method"] == "WEB"
    assert "403" in (result["note"] or "")


@pytest.mark.asyncio
async def test_real_issuer_invalid_from_server_becomes_invalid(monkeypatch):
    monkeypatch.setattr(mock_issuer_adapter, "verify", _mock_boom)
    monkeypatch.setattr(playwright_issuer_adapter, "verify", _playwright_boom)

    async def fake_web(**kwargs):
        return _web_result(
            IssuerVerificationStatus.INVALID,
            error_message="Certificate page does not exist at issuer (HTTP 404)",
        )

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", fake_web)

    issuer = Issuer(
        name="Simplilearn",
        official_domain="simplilearn.com",
        verification_type="WEB",
        active=True,
    )
    result = await _resolve_issuer_record(
        issuer, "SL-9", "https://www.simplilearn.com/cert/SL-9"
    )
    assert result["status"] == "INVALID"
    assert result["method"] == "WEB"


@pytest.mark.asyncio
async def test_demo_issuer_still_resolves_through_mock_registry(monkeypatch):
    async def web_boom(*args, **kwargs):
        raise AssertionError("web fetch queried for a demo issuer")

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", web_boom)

    async def pw_unavailable(**kwargs):
        return _web_result(
            IssuerVerificationStatus.UNAVAILABLE, error_message="no browser"
        )

    monkeypatch.setattr(playwright_issuer_adapter, "verify", pw_unavailable)

    async def mock_valid(**kwargs):
        return _web_result(
            IssuerVerificationStatus.VALID,
            certificate_id_returned="CERT-1001",
            recipient_returned="Rahul Kumar",
            course_returned="Web Security Specialization",
        )

    monkeypatch.setattr(mock_issuer_adapter, "verify", mock_valid)

    issuer = Issuer(
        name="Example University",
        official_domain="example.edu",
        verification_type="WEB",
        verification_url="http://localhost:8001/verify",
        active=True,
    )
    result = await _resolve_issuer_record(
        issuer, "CERT-1001", "https://example.edu/verify?id=CERT-1001"
    )
    assert result["status"] == "VALID"
    assert result["method"] == "WEB->MOCK"
    assert result["recipient"] == "Rahul Kumar"


@pytest.mark.asyncio
async def test_adapter_crash_degrades_to_unavailable(monkeypatch):
    async def raising(**kwargs):
        raise RuntimeError("adapter exploded")

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", raising)

    issuer = Issuer(
        name="MyGreatLearning",
        official_domain="mygreatlearning.com",
        verification_type="WEB",
        active=True,
    )
    result = await _resolve_issuer_record(
        issuer, "MGL-001", "https://www.mygreatlearning.com/x"
    )
    assert result["status"] == "UNAVAILABLE"
    assert "Verification step failed: RuntimeError: adapter exploded" == (
        result["note"] or ""
    )


@pytest.mark.asyncio
async def test_real_issuer_fetches_even_without_certificate_id(monkeypatch):
    monkeypatch.setattr(mock_issuer_adapter, "verify", _mock_boom)
    monkeypatch.setattr(playwright_issuer_adapter, "verify", _playwright_boom)

    async def fake_web(**kwargs):
        return _web_result(
            IssuerVerificationStatus.VALID,
            recipient_returned="Anish Jaiswal",
            course_returned="Programming Fundamentals",
        )

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", fake_web)

    issuer = Issuer(
        name="Simplilearn SkillUp",
        official_domain="app.link",
        verification_type="WEB",
        active=True,
    )
    result = await _resolve_issuer_record(issuer, "", "https://simpli.app.link/abc")
    assert result["status"] == "VALID"
    assert result["method"] == "WEB"
    assert result["recipient"] == "Anish Jaiswal"


@pytest.mark.asyncio
async def test_registry_issuers_still_require_certificate_id(monkeypatch):
    async def web_boom(*args, **kwargs):
        raise AssertionError("web fetch queried for a demo issuer")

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", web_boom)

    issuer = Issuer(
        name="Example University",
        official_domain="example.edu",
        verification_type="WEB",
        verification_url="http://localhost:8001/verify",
        active=True,
    )
    result = await _resolve_issuer_record(issuer, "", "https://example.edu/verify")
    assert result["status"] == "UNAVAILABLE"
    assert result["method"] == "WEB"


# ---------------------------------------------------------------------------
# End-to-end: real issuer rows through the batch API
# ---------------------------------------------------------------------------

MGL_ROWS = [
    [
        "Rahul Kumar",
        "MGL-001",
        "Web Security Specialization",
        "MyGreatLearning",
        "https://www.mygreatlearning.com/certificates/mgl001",
        "",
        "",
    ],
    [
        "Sarah Chen",
        "MGL-002",
        "",
        "MyGreatLearning",
        "https://www.mygreatlearning.com/certificates/mgl002",
        "",
        "",
    ],
    [
        "Unknown Student",
        "MGL-003",
        "",
        "MyGreatLearning",
        "https://www.mygreatlearning.com/certificates/mgl003",
        "",
        "",
    ],
]


def _make_workbook(rows):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(list(ALL_COLUMNS))
    for row in rows:
        ws.append(list(row))
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


async def _ensure_mgl_issuer(db_session):
    from sqlalchemy import select

    existing = await db_session.execute(
        select(Issuer).where(Issuer.official_domain == "mygreatlearning.com")
    )
    if existing.scalar_one_or_none() is None:
        db_session.add(
            Issuer(
                name="MyGreatLearning",
                official_domain="mygreatlearning.com",
                verification_type="WEB",
                active=True,
            )
        )
        await db_session.commit()


@pytest.mark.asyncio
async def test_real_issuer_rows_produce_mixed_verdicts(
    student_auth_headers, db_session, monkeypatch
):
    await _ensure_mgl_issuer(db_session)

    async def fake_web(**kwargs):
        cid = kwargs.get("certificate_id")
        if cid == "MGL-001":
            return _web_result(
                IssuerVerificationStatus.VALID,
                certificate_id_returned="MGL-001",
                recipient_returned="Rahul Kumar",
                course_returned="Web Security Specialization",
                status_returned="VALID",
                raw_evidence='{"http_status": 200}',
            )
        if cid == "MGL-002":
            return _web_result(
                IssuerVerificationStatus.INVALID,
                certificate_id_returned="MGL-002",
                error_message="Certificate page does not exist at issuer (HTTP 404)",
            )
        return _web_result(
            IssuerVerificationStatus.UNAVAILABLE,
            certificate_id_returned="MGL-003",
            error_message="Issuer site returned HTTP 403 (access restricted or outage)",
        )

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", fake_web)
    monkeypatch.setattr(mock_issuer_adapter, "verify", _mock_boom)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        upload = await client.post(
            "/api/batch",
            files={"file": ("mgl.xlsx", _make_workbook(MGL_ROWS), WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert upload.status_code == 200, upload.text
        batch_id = upload.json()["id"]

        process = await client.post(
            f"/api/batch/{batch_id}/process", headers=student_auth_headers
        )
        assert process.status_code == 200, process.text
        done = process.json()
        assert done["status"] == "COMPLETED"
        counts = done["verdict_counts"]
        assert counts[VERDICT_LEGIT] == 1
        assert counts[VERDICT_FAKE] == 1
        assert counts[VERDICT_ANOMALY] == 1

        result = await client.get(
            f"/api/batch/{batch_id}/result", headers=student_auth_headers
        )
        assert result.status_code == 200
        wb = load_workbook(io.BytesIO(result.content))
        ws = wb[SHEET_NAME]
        verdict_col = len(INPUT_COLUMNS) + 1
        analysis_col = len(INPUT_COLUMNS) + 4
        verdicts = [ws.cell(row=r, column=verdict_col).value for r in range(2, 5)]
        assert verdicts == [VERDICT_LEGIT, VERDICT_FAKE, VERDICT_ANOMALY]
        # The UNAVAILABLE row must surface why the issuer check failed.
        assert "Issuer check:" in (ws.cell(row=4, column=analysis_col).value or "")


@pytest.mark.asyncio
async def test_crashing_issuer_check_still_completes_batch(
    student_auth_headers, db_session, monkeypatch
):
    """Regression: a wrapped URL plus an adapter bug must not fail the batch."""
    await _ensure_mgl_issuer(db_session)

    async def raising(**kwargs):
        raise RuntimeError("adapter exploded")

    monkeypatch.setattr(web_fetch_issuer_adapter, "verify", raising)
    monkeypatch.setattr(mock_issuer_adapter, "verify", _mock_boom)

    rows = [list(row) for row in MGL_ROWS[:1]]
    rows[0][0] = "Test Person"
    rows[0][1] = "MGL-901"
    rows[0][4] = "https://www.mygreatlearning.com/certificates/\nmgl901"

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        upload = await client.post(
            "/api/batch",
            files={"file": ("crash.xlsx", _make_workbook(rows), WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert upload.status_code == 200, upload.text
        batch_id = upload.json()["id"]

        process = await client.post(
            f"/api/batch/{batch_id}/process", headers=student_auth_headers
        )
        assert process.status_code == 200, process.text
        done = process.json()
        assert done["status"] == "COMPLETED"
        counts = done["verdict_counts"]
        assert counts[VERDICT_ANOMALY] == 1
        assert counts[VERDICT_LEGIT] == 0
        assert counts[VERDICT_FAKE] == 0

        result = await client.get(
            f"/api/batch/{batch_id}/result", headers=student_auth_headers
        )
        wb = load_workbook(io.BytesIO(result.content))
        ws = wb[SHEET_NAME]
        analysis = ws.cell(row=2, column=len(INPUT_COLUMNS) + 4).value or ""
        assert "Verification step failed: RuntimeError" in analysis


# ---------------------------------------------------------------------------
# Issuer-configured existence check (Cursa-style object storage)
# ---------------------------------------------------------------------------

CURSA_META = {
    "existence_check": {
        "url_template": "http://localhost:8001/cert_{code}.png",
        "strip_prefix": "cert",
    },
    "allow_localhost": True,
}
CURSA_LIKE_URL = (
    "http://localhost:8001/en/my-certificate/cert18c8acd5281a4a5b5a4f312cd760a25b"
)


@pytest.mark.asyncio
async def test_existence_check_object_present_is_valid(monkeypatch):
    seen = {}

    def responder(url):
        seen["url"] = url
        return _resp(url, 200, "", content_type="image/png")

    _install_fetch(monkeypatch, responder)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="25BBTCS013",
        verification_url=CURSA_LIKE_URL,
        metadata=CURSA_META,
    )
    assert result.verification_status == IssuerVerificationStatus.VALID
    assert result.recipient_returned is None
    assert seen["url"].endswith("cert_18c8acd5281a4a5b5a4f312cd760a25b.png")
    assert "not machine-readable" in (result.error_message or "")


@pytest.mark.asyncio
async def test_existence_check_object_missing_is_invalid(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 404, ""))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="25BBTCS013",
        verification_url=CURSA_LIKE_URL,
        metadata=CURSA_META,
    )
    assert result.verification_status == IssuerVerificationStatus.INVALID
    assert "not found (existence check)" in (result.error_message or "")


@pytest.mark.asyncio
async def test_existence_check_forbidden_is_unavailable_not_invalid(monkeypatch):
    _install_fetch(monkeypatch, lambda url: _resp(url, 403, ""))
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="25BBTCS013",
        verification_url=CURSA_LIKE_URL,
        metadata=CURSA_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE


@pytest.mark.asyncio
async def test_existence_check_without_strip_prefix_uses_full_segment(monkeypatch):
    meta = {
        "existence_check": {"url_template": "http://localhost:8001/{code}.png"},
        "allow_localhost": True,
    }
    seen = {}

    def responder(url):
        seen["url"] = url
        return _resp(url, 404, "")

    _install_fetch(monkeypatch, responder)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="X",
        verification_url=CURSA_LIKE_URL,
        metadata=meta,
    )
    assert result.verification_status == IssuerVerificationStatus.INVALID
    assert "cert18c8acd5281a4a5b5a4f312cd760a25b.png" in seen["url"]
    assert "cert_cert" not in seen["url"]


@pytest.mark.asyncio
async def test_existence_check_url_without_code_is_unavailable(monkeypatch):
    _install_fetch_boom(monkeypatch)
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="X",
        verification_url="http://localhost:8001/",
        metadata=CURSA_META,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "no certificate code" in (result.error_message or "")


@pytest.mark.asyncio
async def test_existence_check_misconfigured_template_is_unavailable(monkeypatch):
    _install_fetch_boom(monkeypatch)
    meta = {
        "existence_check": {"url_template": "http://localhost:8001/{other}.png"},
        "allow_localhost": True,
    }
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="X",
        verification_url=CURSA_LIKE_URL,
        metadata=meta,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "misconfigured" in (result.error_message or "")


@pytest.mark.asyncio
async def test_existence_check_ssrf_target_is_unavailable(monkeypatch):
    _install_fetch_boom(monkeypatch)
    meta = {
        "existence_check": {"url_template": "http://169.254.169.254/{code}.png"},
        "allow_localhost": True,
    }
    result = await web_fetch_issuer_adapter.verify(
        certificate_id="X",
        verification_url=CURSA_LIKE_URL,
        metadata=meta,
    )
    assert result.verification_status == IssuerVerificationStatus.UNAVAILABLE
    assert "SSRF" in (result.error_message or "")
