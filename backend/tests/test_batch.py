import io

import pytest
from httpx import ASGITransport, AsyncClient
from openpyxl import load_workbook

from app.batch.template import (
    ALL_COLUMNS,
    INPUT_COLUMNS,
    OUTPUT_COLUMNS,
    SHEET_NAME,
    VERDICT_ANOMALY,
    VERDICT_FAKE,
    VERDICT_LEGIT,
    WorkbookFormatError,
    build_template_bytes,
    parse_workbook,
    write_results,
)
from app.main import app

WORKBOOK_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def make_workbook(rows, sheet_name=SHEET_NAME, headers=None, extra_headers=None):
    """Builds an in-memory workbook following (or breaking) the format contract."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    ws.append(headers if headers is not None else ALL_COLUMNS)
    for row in rows:
        ws.append(list(row))
    if extra_headers:
        base = len(ALL_COLUMNS)
        for offset, name in enumerate(extra_headers):
            ws.cell(row=1, column=base + 1 + offset, value=name)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


VALID_ROWS = [
    # LEGIT: issuer registry confirms recipient and course
    [
        "Rahul Kumar",
        "CERT-1001",
        "Web Security Specialization",
        "Example University",
        "https://example.edu/verify?id=CERT-1001",
        "2024-05-01",
        "",
    ],
    # FAKE: issuer declares this ID revoked / non-existent
    [
        "Sarah Chen",
        "CERT-1003",
        "Cloud Architecture",
        "Example University",
        "https://example.edu/verify?id=CERT-1003",
        "2024-01-01",
        "",
    ],
    # ANOMALY: issuer returns a different recipient than the sheet claims
    [
        "Someone Else",
        "CERT-1002",
        "Machine Learning",
        "Example University",
        "https://example.edu/verify?id=CERT-1002",
        "",
        "",
    ],
    # ANOMALY: duplicate certificate ID inside the same sheet
    [
        "Rahul Kumar",
        "CERT-1001",
        "Web Security Specialization",
        "Example University",
        "https://example.edu/verify?id=CERT-1001",
        "2024-05-01",
        "",
    ],
]


# ---------------------------------------------------------------------------
# Format contract
# ---------------------------------------------------------------------------


def test_template_has_locked_analysis_zone():
    content = build_template_bytes()
    assert content.startswith(b"PK\x03\x04")

    wb = load_workbook(io.BytesIO(content))
    ws = wb[SHEET_NAME]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 13)]
    assert headers == ALL_COLUMNS
    assert ws.protection.sheet is True
    # Header and output/analysis cells locked, input cells editable.
    assert ws.cell(row=1, column=1).protection.locked is True
    assert ws.cell(row=2, column=1).protection.locked is False
    assert ws.cell(row=2, column=len(INPUT_COLUMNS) + 1).protection.locked is True
    assert "FORMAT_GUIDE" in wb.sheetnames


def test_parse_accepts_conforming_workbook():
    rows = parse_workbook(make_workbook(VALID_ROWS))
    assert len(rows) == 4
    assert rows[0].recipient_name == "Rahul Kumar"
    assert rows[0].certificate_id == "CERT-1001"


def test_parse_rejects_wrong_sheet_name():
    with pytest.raises(WorkbookFormatError) as exc:
        parse_workbook(make_workbook(VALID_ROWS, sheet_name="Sheet1"))
    assert "Certificates" in str(exc.value)


def test_parse_rejects_reordered_headers():
    headers = list(ALL_COLUMNS)
    headers[0], headers[1] = headers[1], headers[0]
    with pytest.raises(WorkbookFormatError) as exc:
        parse_workbook(make_workbook(VALID_ROWS, headers=headers))
    assert "Header row must be exactly" in str(exc.value)


def test_parse_rejects_missing_and_unknown_columns():
    with pytest.raises(WorkbookFormatError) as exc:
        parse_workbook(make_workbook(VALID_ROWS, headers=INPUT_COLUMNS))
    assert "Missing" in str(exc.value)
    assert "VERDICT" in str(exc.value)  # output zone missing

    with pytest.raises(WorkbookFormatError) as exc2:
        parse_workbook(make_workbook(VALID_ROWS, extra_headers=["SMUGGLED_VERDICT"]))
    assert "SMUGGLED_VERDICT" in str(exc2.value)


def test_parse_rejects_prefilled_verdict_column():
    """A client must never be able to smuggle its own verdict into the sheet."""
    rows = [
        list(r) + ["LEGIT", "NONE", "NONE", "pre-written", "0.99"]
        for r in VALID_ROWS[:1]
    ]
    with pytest.raises(WorkbookFormatError) as exc:
        parse_workbook(make_workbook(rows))
    assert "reserved output column" in str(exc.value)
    assert "VERDICT" in str(exc.value)


def test_parse_rejects_empty_workbook():
    with pytest.raises(WorkbookFormatError) as exc:
        parse_workbook(make_workbook([]))
    assert "no data rows" in str(exc.value)


def test_write_results_preserves_input_and_locks_output():
    content = make_workbook(VALID_ROWS)
    results = [
        {"row_number": 2, "verdict": VERDICT_LEGIT},
        {"row_number": 3, "verdict": VERDICT_FAKE},
        {"row_number": 4, "verdict": VERDICT_ANOMALY},
    ]
    from app.batch.template import RowVerdict

    annotated = write_results(
        content,
        [
            RowVerdict(r["row_number"], r["verdict"], "NONE", "NONE", "ok", 0.9)
            for r in results
        ],
    )

    wb = load_workbook(io.BytesIO(annotated))
    ws = wb[SHEET_NAME]
    verdict_col = len(INPUT_COLUMNS) + 1

    assert [ws.cell(row=i, column=verdict_col).value for i in (2, 3, 4)] == [
        VERDICT_LEGIT,
        VERDICT_FAKE,
        VERDICT_ANOMALY,
    ]
    # Input cells untouched
    assert ws.cell(row=2, column=1).value == "Rahul Kumar"
    assert ws.cell(row=2, column=2).value == "CERT-1001"
    # Zones: input editable, output locked, sheet protected
    assert ws.cell(row=2, column=1).protection.locked is False
    assert ws.cell(row=2, column=verdict_col).protection.locked is True
    assert ws.protection.sheet is True
    # Output headers intact
    assert [
        ws.cell(row=1, column=c).value
        for c in range(verdict_col, verdict_col + len(OUTPUT_COLUMNS))
    ] == OUTPUT_COLUMNS


# ---------------------------------------------------------------------------
# API flow
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_template_endpoint_requires_auth_and_returns_workbook(
    student_auth_headers,
):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        anon = await client.get("/api/batch/template")
        assert anon.status_code == 401

        authed = await client.get("/api/batch/template", headers=student_auth_headers)
        assert authed.status_code == 200
        assert authed.content.startswith(b"PK\x03\x04")
        wb = load_workbook(io.BytesIO(authed.content))
        assert SHEET_NAME in wb.sheetnames
        assert "FORMAT_GUIDE" in wb.sheetnames


@pytest.mark.asyncio
async def test_upload_rejects_non_workbook(student_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/batch",
            files={"file": ("certs.csv", b"a,b,c\n1,2,3", "text/csv")},
            headers=student_auth_headers,
        )
        assert resp.status_code == 400
        assert "INVALID_EXTENSION" in resp.json()["detail"]


@pytest.mark.asyncio
async def test_upload_rejects_format_violations(student_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # Wrong sheet name
        bad_sheet = make_workbook(VALID_ROWS, sheet_name="Sheet1")
        resp = await client.post(
            "/api/batch",
            files={"file": ("bad.xlsx", bad_sheet, WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert resp.status_code == 400
        detail = resp.json()["detail"]
        assert detail["message"] == "Workbook does not follow the required format"
        assert any("Certificates" in e for e in detail["errors"])

        # Pre-filled analysis zone
        prefill = make_workbook(
            [
                list(r) + ["LEGIT", "NONE", "NONE", "forged", "0.99"]
                for r in VALID_ROWS[:1]
            ]
        )
        resp2 = await client.post(
            "/api/batch",
            files={"file": ("forged.xlsx", prefill, WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert resp2.status_code == 400
        assert any(
            "reserved output column" in e for e in resp2.json()["detail"]["errors"]
        )


@pytest.mark.asyncio
async def test_full_batch_analysis_flow(student_auth_headers, teacher_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Upload a conforming workbook
        upload = await client.post(
            "/api/batch",
            files={"file": ("certs.xlsx", make_workbook(VALID_ROWS), WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert upload.status_code == 200, upload.text
        batch = upload.json()
        batch_id = batch["id"]
        assert batch["total_rows"] == 4

        # 2. Run the analysis deterministically
        process = await client.post(
            f"/api/batch/{batch_id}/process", headers=student_auth_headers
        )
        assert process.status_code == 200, process.text
        done = process.json()
        assert done["status"] == "COMPLETED"
        assert done["processed_rows"] == 4
        assert done["result_ready"] is True

        counts = done["verdict_counts"]
        assert counts[VERDICT_LEGIT] == 1
        assert counts[VERDICT_FAKE] == 1
        assert counts[VERDICT_ANOMALY] == 2

        # 3. Download annotated workbook
        result = await client.get(
            f"/api/batch/{batch_id}/result", headers=student_auth_headers
        )
        assert result.status_code == 200
        assert result.content.startswith(b"PK\x03\x04")

        wb = load_workbook(io.BytesIO(result.content))
        ws = wb[SHEET_NAME]
        verdict_col = len(INPUT_COLUMNS) + 1
        verdicts = [ws.cell(row=r, column=verdict_col).value for r in range(2, 6)]
        assert verdicts == [
            VERDICT_LEGIT,
            VERDICT_FAKE,
            VERDICT_ANOMALY,
            VERDICT_ANOMALY,
        ]
        codes = [ws.cell(row=r, column=verdict_col + 2).value for r in range(2, 6)]
        assert "A006" in (codes[1] or "")  # revoked / non-existent ID
        assert "A007" in (codes[2] or "")  # recipient differs from issuer record
        assert ws.protection.sheet is True

        # 4. Input zone still intact
        assert ws.cell(row=2, column=1).value == "Rahul Kumar"

        # 5. Listing includes the batch for the teacher too
        listing = await client.get("/api/batch", headers=teacher_auth_headers)
        assert listing.status_code == 200
        assert any(b["id"] == batch_id for b in listing.json())


@pytest.mark.asyncio
async def test_student_cannot_access_other_students_batch(student_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        upload = await client.post(
            "/api/batch",
            files={"file": ("mine.xlsx", make_workbook(VALID_ROWS), WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert upload.status_code == 200
        batch_id = upload.json()["id"]

        from app.security.auth import create_access_token

        other_headers = {
            "Authorization": f"Bearer {create_access_token(data={'sub': 'test-user-student-2', 'role': 'STUDENT'})}"
        }
        resp = await client.get(f"/api/batch/{batch_id}", headers=other_headers)
        assert resp.status_code == 403

        dl = await client.get(f"/api/batch/{batch_id}/result", headers=other_headers)
        assert dl.status_code == 403

        listing = await client.get("/api/batch", headers=other_headers)
        assert listing.status_code == 200
        assert all(b["id"] != batch_id for b in listing.json())


@pytest.mark.asyncio
async def test_result_not_ready_returns_conflict(student_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        upload = await client.post(
            "/api/batch",
            files={"file": ("pending.xlsx", make_workbook(VALID_ROWS), WORKBOOK_MIME)},
            headers=student_auth_headers,
        )
        assert upload.status_code == 200
        batch_id = upload.json()["id"]

        # Background task has completed by now; force reset to PROCESSING to
        # assert the guard for genuinely pending batches.
        from app.database import AsyncSessionLocal
        from app.models import BatchAnalysis, BatchStatus

        async with AsyncSessionLocal() as session:
            from sqlalchemy import select

            res = await session.execute(
                select(BatchAnalysis).where(BatchAnalysis.id == batch_id)
            )
            batch = res.scalar_one()
            batch.status = BatchStatus.PROCESSING.value
            batch.result_storage_key = None
            await session.commit()

        resp = await client.get(
            f"/api/batch/{batch_id}/result", headers=student_auth_headers
        )
        assert resp.status_code == 409


# ---------------------------------------------------------------------------
# Google Sheets intake
# ---------------------------------------------------------------------------


def make_google_sheet_bytes(with_verdict_column: bool = False):
    """A messy, title-row-first sheet in the shape of a real class spreadsheet."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["CERTIFICATE VERIFICATION REPORT - 2024"])
    headers = [
        "SL. NO.",
        "NAME OF THE STUDENT",
        "",
        "REGISTRATION  NUMBER",
        "LINKS OF CERTIFICATES",
    ]
    if with_verdict_column:
        headers.append("VERDICT")
    ws.append(headers)
    rows = [
        [1, "Rahul Kumar", "", "CERT-1001", "https://example.edu/verify?id=CERT-1001"],
        [2, "Fraud Person", "", "CERT-1003", "https://example.edu/verify?id=CERT-1003"],
    ]
    for row in rows:
        if with_verdict_column:
            row = row + ["LEGIT"]
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class TestExtractSheetId:
    def test_accepts_common_share_links(self):
        from app.batch.google_sheets import extract_sheet_id

        sheet_id = "1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW"
        for url in (
            f"https://docs.google.com/spreadsheets/d/{sheet_id}/edit?usp=sharing",
            f"https://docs.google.com/spreadsheets/d/{sheet_id}/view",
            f"http://docs.google.com/spreadsheets/d/{sheet_id}/export?format=xlsx",
            f"https://sheets.google.com/spreadsheets/d/{sheet_id}/edit#gid=0",
            f"  https://docs.google.com/spreadsheets/d/{sheet_id}/edit  ",
        ):
            assert extract_sheet_id(url) == sheet_id

    def test_rejects_non_google_hosts(self):
        from app.batch.google_sheets import GoogleSheetError, extract_sheet_id

        sheet_id = "1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW"
        for url in (
            f"https://docs.google.com.evil.com/spreadsheets/d/{sheet_id}/edit",
            f"https://evil.com/spreadsheets/d/{sheet_id}/edit",
            f"https://notsheets.example.org/spreadsheets/d/{sheet_id}/edit",
        ):
            with pytest.raises(GoogleSheetError, match="docs.google.com"):
                extract_sheet_id(url)

    def test_rejects_bad_scheme_and_missing_id(self):
        from app.batch.google_sheets import GoogleSheetError, extract_sheet_id

        with pytest.raises(GoogleSheetError, match="https://"):
            extract_sheet_id("javascript:alert(1)")
        with pytest.raises(GoogleSheetError, match="https://"):
            extract_sheet_id("ftp://docs.google.com/spreadsheets/d/abcdef123456/edit")
        with pytest.raises(GoogleSheetError, match="spreadsheet ID"):
            extract_sheet_id("https://docs.google.com/spreadsheets/d/abc/edit")
        with pytest.raises(GoogleSheetError, match="spreadsheet ID"):
            extract_sheet_id("https://docs.google.com/spreadsheets/u/0/")
        with pytest.raises(GoogleSheetError, match="https://"):
            extract_sheet_id("")


class TestAdaptWorkbook:
    def test_contract_workbook_passes_through(self):
        from app.batch.google_sheets import adapt_workbook

        content = make_workbook(VALID_ROWS)
        adapted, mode = adapt_workbook(content)
        assert mode == "contract"
        assert adapted == content

    def test_maps_messy_headers_and_skips_title_row(self):
        from app.batch.google_sheets import adapt_workbook

        adapted, mode = adapt_workbook(make_google_sheet_bytes())
        assert mode == "mapped"

        rows = parse_workbook(adapted)  # adapted output must satisfy the contract
        assert len(rows) == 2
        assert rows[0].recipient_name == "Rahul Kumar"
        assert rows[0].certificate_id == "CERT-1001"
        assert rows[0].certificate_url == "https://example.edu/verify?id=CERT-1001"
        assert rows[1].recipient_name == "Fraud Person"
        # Unmappable columns (SL. NO., blank header) are simply dropped.

    def test_source_verdict_column_never_reaches_output_zone(self):
        """Even if the remote sheet ships its own verdicts, the output zone stays empty."""
        from app.batch.google_sheets import adapt_workbook

        adapted, mode = adapt_workbook(
            make_google_sheet_bytes(with_verdict_column=True)
        )
        assert mode == "mapped"

        wb = load_workbook(io.BytesIO(adapted))
        ws = wb[SHEET_NAME]
        verdict_col = len(INPUT_COLUMNS) + 1
        assert [ws.cell(row=r, column=verdict_col).value for r in (2, 3)] == [
            None,
            None,
        ]

    def test_rejects_unmappable_sheet(self):
        from app.batch.google_sheets import adapt_workbook

        unmappable = make_workbook(
            [["a", "b"], ["1", "2"]],
            sheet_name="Sheet1",
            headers=["foo", "bar"],
        )
        with pytest.raises(WorkbookFormatError, match="usable header row"):
            adapt_workbook(unmappable)

    def test_rejects_garbage_bytes(self):
        from app.batch.google_sheets import adapt_workbook

        with pytest.raises(WorkbookFormatError, match="could not be opened"):
            adapt_workbook(b"not a workbook at all")


@pytest.mark.asyncio
async def test_from_url_rejects_invalid_urls_without_fetching(student_auth_headers):
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        anon = await client.post(
            "/api/batch/from-url", json={"url": "https://docs.google.com/x"}
        )
        assert anon.status_code == 401

        for url in (
            "https://evil.com/spreadsheets/d/1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW/edit",
            "https://docs.google.com.evil.com/spreadsheets/d/1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW/edit",
            "javascript:alert(1)",
            "https://docs.google.com/spreadsheets/d/abc/edit",
        ):
            resp = await client.post(
                "/api/batch/from-url", json={"url": url}, headers=student_auth_headers
            )
            assert resp.status_code == 400, url
            assert isinstance(resp.json()["detail"], str)


@pytest.mark.asyncio
async def test_from_url_end_to_end(student_auth_headers, monkeypatch):
    """URL -> fetch (mocked) -> adapt -> ingest -> analyse -> verdicts."""

    async def fake_fetch(sheet_id):
        assert sheet_id == "1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW"
        return make_google_sheet_bytes()

    monkeypatch.setattr("app.api.batch.fetch_sheet_export", fake_fetch)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/batch/from-url",
            json={
                "url": "https://docs.google.com/spreadsheets/d/1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW/edit?usp=sharing"
            },
            headers=student_auth_headers,
        )
        assert resp.status_code == 200, resp.text
        batch = resp.json()
        assert batch["total_rows"] == 2
        assert batch["original_filename"].startswith("gsheet_")

        process = await client.post(
            f"/api/batch/{batch['id']}/process", headers=student_auth_headers
        )
        assert process.status_code == 200, process.text
        done = process.json()
        assert done["status"] == "COMPLETED"
        counts = done["verdict_counts"]
        assert counts[VERDICT_LEGIT] == 1  # CERT-1001 confirmed by issuer registry
        assert counts[VERDICT_FAKE] == 1  # CERT-1003 declared non-existent
        assert counts.get(VERDICT_ANOMALY, 0) == 0


@pytest.mark.asyncio
async def test_from_url_reports_adaptation_errors(student_auth_headers, monkeypatch):
    async def fake_fetch(sheet_id):
        return make_workbook(
            [["a", "b"], ["1", "2"]], sheet_name="Sheet1", headers=["foo", "bar"]
        )

    monkeypatch.setattr("app.api.batch.fetch_sheet_export", fake_fetch)

    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        resp = await client.post(
            "/api/batch/from-url",
            json={
                "url": "https://docs.google.com/spreadsheets/d/1l_CoBQ6BnSTDADpCBX6ibcaAIapiFyTW/edit"
            },
            headers=student_auth_headers,
        )
        assert resp.status_code == 400
        detail = resp.json()["detail"]
        assert detail["message"] == "Workbook does not follow the required format"
        assert any("usable header row" in e for e in detail["errors"])
